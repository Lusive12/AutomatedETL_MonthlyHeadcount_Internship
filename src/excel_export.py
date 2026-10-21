"""
excel_export.py — Shared Excel writer with safe overwrite and date formatting.

Handles the common PermissionError that occurs when Excel has the
output file open. Gives the user a clear actionable error message.
Supports applying Excel Date number formatting ('DD/MM/YYYY') to date columns.
"""
import logging
from pathlib import Path
from typing import List, Optional
import openpyxl
import pandas as pd


def export_df(
    df: pd.DataFrame,
    path: Path,
    logger: logging.Logger = None,
    date_columns: Optional[List[str]] = None,
) -> None:
    """
    Write DataFrame to Excel.
    - Auto-creates parent directories.
    - Deletes any existing file before writing to avoid lock errors.
    - If date_columns provided, sets Excel cell number_format to 'DD/MM/YYYY'.
    """
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if output_path.exists():
        try:
            output_path.unlink()
        except PermissionError:
            msg = (
                f"Cannot overwrite '{output_path.name}' — "
                f"the file is open in Excel. Close it and run again."
            )
            if logger:
                logger.error(f"[Export] {msg}")
            raise PermissionError(msg)

    df.to_excel(output_path, index=False)

    # Format specified date columns to Excel Date format 'DD/MM/YYYY'
    if date_columns and output_path.exists():
        try:
            wb = openpyxl.load_workbook(output_path)
            ws = wb.active

            header_map = {}
            for col_idx in range(1, ws.max_column + 1):
                val = ws.cell(row=1, column=col_idx).value
                if val is not None:
                    header_map[str(val).strip().lower()] = col_idx

            matched_indices = []
            for dcol in date_columns:
                target = str(dcol).strip().lower()
                if target in header_map:
                    matched_indices.append(header_map[target])

            for col_idx in matched_indices:
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if cell.value is not None:
                        cell.number_format = "DD/MM/YYYY"

            wb.save(output_path)
            wb.close()
            if logger:
                logger.info(f"[Export] Applied Excel Date number format (DD/MM/YYYY) to {len(matched_indices)} column(s).")
        except Exception as e:
            if logger:
                logger.warning(f"[Export] Could not apply Date formatting via openpyxl: {e}")

    if logger:
        logger.info(f"[Export] Saved: {output_path.name} ({len(df):,} rows)")
